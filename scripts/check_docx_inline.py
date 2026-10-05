"""End-to-end probe: the real analysis report exported to .docx must contain
no literal ** or ` inline-markdown characters in any paragraph or table cell."""

import asyncio
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document  # noqa: E402

from agent.agent import run_analysis  # noqa: E402
from agent.docx_export import markdown_to_docx  # noqa: E402
from agent.synth import make_income  # noqa: E402


def main() -> int:
    artifacts = tempfile.mkdtemp(prefix="daa_probe_docx_")
    ws = asyncio.run(
        run_analysis(make_income(), "男性和女性的 income 是否存在显著差异？", artifacts_dir=artifacts)
    )
    assert "**[" in ws.report_md, "fixture report should contain inline bold markers"

    out = Path(artifacts) / "report.docx"
    markdown_to_docx(ws.report_md, out, image_dir=artifacts)

    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs]
    texts += [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    joined = "\n".join(texts)
    bad_star = sum("*" in t for t in texts)
    bad_tick = sum("`" in t for t in texts)
    bold_runs = [r for p in doc.paragraphs for r in p.runs if r.bold]
    print(f"paragraph+cell texts: {len(texts)}, with '*': {bad_star}, with '`': {bad_tick}")
    print(f"bold runs: {len(bold_runs)}, e.g. {[r.text for r in bold_runs[:3]]}")
    ok = bad_star == 0 and bad_tick == 0 and bold_runs
    print("PASS" if ok else "FAIL")
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
