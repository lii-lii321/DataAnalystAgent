import asyncio
from pathlib import Path

from agent.agent import run_analysis
from agent.docx_export import markdown_to_docx
from agent.synth import make_income


def test_markdown_report_to_docx_with_charts(tmp_path):
    ws = asyncio.run(
        run_analysis(
            make_income(),
            "男性和女性的 income 是否存在显著差异？",
            artifacts_dir=str(tmp_path),
        )
    )
    out = tmp_path / "report.docx"
    markdown_to_docx(ws.report_md, out, image_dir=tmp_path)

    data = out.read_bytes()
    assert data[:2] == b"PK"
    assert len(data) > 5000


def test_docx_renders_tables_and_headings(tmp_path):
    md = "# 报告\n\n## 统计检验\n\n| 检验 | p |\n| --- | --- |\n| t-test | 0.01 |\n\n- 发现一\n- 发现二\n"
    out = tmp_path / "mini.docx"
    markdown_to_docx(md, out)

    from docx import Document

    doc = Document(str(out))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith(("Title", "Heading"))]
    assert "报告" in headings
    assert "统计检验" in headings
    assert len(doc.tables) == 1
    assert doc.tables[0].cell(1, 0).text == "t-test"
    bullets = [p.text for p in doc.paragraphs if p.style.name == "List Bullet"]
    assert bullets == ["发现一", "发现二"]


def test_docx_strips_inline_markdown(tmp_path):
    md = (
        "# 报告\n\n"
        "- **[high]** income：缺失 12 个值（建议：中位数填充）\n"
        "- 任务：回归，目标列：`income`（train=800, test=200）\n\n"
        "| 列 | 缺失 |\n"
        "| --- | --- |\n"
        "| **[medium]** age | 3 |\n"
    )
    out = tmp_path / "inline.docx"
    markdown_to_docx(md, out)

    from docx import Document

    doc = Document(str(out))
    joined = "\n".join(p.text for p in doc.paragraphs)
    assert "*" not in joined
    assert "`" not in joined

    bold_runs = [run for p in doc.paragraphs for run in p.runs if run.bold]
    assert bold_runs
    assert any("[high]" in run.text for run in bold_runs)

    mono_runs = [run for p in doc.paragraphs for run in p.runs if run.font.name == "Consolas"]
    assert any(run.text == "income" for run in mono_runs)

    cell_text = doc.tables[0].cell(1, 0).text
    assert cell_text == "[medium] age"
    assert "*" not in cell_text
