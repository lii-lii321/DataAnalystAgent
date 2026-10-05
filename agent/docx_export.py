import re
from pathlib import Path

from docx import Document
from docx.shared import Inches


def markdown_to_docx(md: str, out_path, image_dir=None) -> str:
    """Render the analysis report (headings, bullets, tables, images) into a .docx file."""
    doc = Document()
    lines = md.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("|"):
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                if not all(set(cell) <= set("-: ") for cell in cells):
                    rows.append(cells)
                index += 1
            if rows:
                width = max(len(row) for row in rows)
                table = doc.add_table(rows=len(rows), cols=width)
                table.style = "Table Grid"
                for row_no, row in enumerate(rows):
                    for col_no in range(width):
                        table.cell(row_no, col_no).text = row[col_no] if col_no < len(row) else ""
            continue

        if stripped.startswith("![") and "](" in stripped:
            match = re.search(r"\]\(([^)]+)\)", stripped)
            if match:
                candidate = Path(image_dir or ".") / Path(match.group(1)).name
                if candidate.exists():
                    doc.add_picture(str(candidate), width=Inches(5.8))
                index += 1
                continue

        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=2)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=1)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=0)
        elif stripped.startswith("- "):
            doc.add_paragraph(stripped[2:], style="List Bullet")
        elif stripped:
            doc.add_paragraph(stripped)
        index += 1

    doc.save(out_path)
    return str(out_path)
